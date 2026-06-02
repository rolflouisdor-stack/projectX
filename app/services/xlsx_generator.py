"""Generate the downloadable .xlsx files for completed jobs.

Result files are written to DigitalOcean Spaces (S3-compatible object
storage) rather than the local filesystem so they survive App Platform
container restarts. The API hands out short-lived presigned download URLs.

Scrub artifacts read the real rows out of scrub_job_records + their EAV
custom fields. Purchase artifacts are still a stub generator since
purchase_jobs has no underlying record table yet.
"""
import io
import random
import string
import logging
from datetime import datetime, timedelta
from openpyxl import Workbook

from app.extensions import get_db
from app.models.scrub_job_record import ScrubJobRecord
from app.models.scrub_job_record_field import ScrubJobRecordField
from app.models.scrub_job_field_mapping import ScrubJobFieldMapping, STANDARD_FIELDS
from app.services import storage

logger = logging.getLogger(__name__)


def _rand_email():
    user = ''.join(random.choices(string.ascii_lowercase, k=random.randint(5, 10)))
    domain = random.choice(['gmail.com', 'yahoo.com', 'outlook.com', 'hotmail.com', 'aol.com'])
    return f'{user}@{domain}'


def _workbook_to_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_scrub_artifact(job):
    """Produce an .xlsx of the unique-and-valid records from a scrub job.

    Reads `scrub_job_records` rows where is_unique=true AND is_valid=true,
    pulls the matching EAV custom fields, and emits a workbook whose column
    set is (standard fields the user mapped) + (custom fields the user
    mapped). Skipped columns are omitted.

    Returns (filename, s3_key). The caller stores both on the job row.
    """
    db = get_db()

    mappings = (db.query(ScrubJobFieldMapping)
                .filter_by(scrub_job_id=job.id, skip=False)
                .order_by(ScrubJobFieldMapping.column_index)
                .all())
    if not mappings:
        raise RuntimeError(f"scrub_job {job.id} has no field mappings — cannot build artifact")

    standard_targets = [m.target_field for m in mappings if m.is_standard and m.target_field in STANDARD_FIELDS]
    custom_targets = [m.target_field for m in mappings if not m.is_standard]
    headers = standard_targets + custom_targets

    wb = Workbook(write_only=True)
    ws = wb.create_sheet('records')
    ws.append(headers)

    # Keyset pagination (WHERE id > last_id), NOT OFFSET. OFFSET re-scans and
    # discards `offset` rows on every page → O(n²) on large result sets, which
    # blew the request timeout on a 138k-row job. Keyset is linear.
    BATCH = 1000
    last_id = 0
    while True:
        recs = (db.query(ScrubJobRecord)
                .filter(ScrubJobRecord.scrub_job_id == job.id,
                        ScrubJobRecord.is_unique.is_(True),
                        ScrubJobRecord.is_valid.is_(True),
                        ScrubJobRecord.id > last_id)
                .order_by(ScrubJobRecord.id)
                .limit(BATCH).all())
        if not recs:
            break
        rec_ids = [r.id for r in recs]
        # Bulk-load custom fields for this batch, group by record id.
        extras = (db.query(ScrubJobRecordField)
                  .filter(ScrubJobRecordField.record_id.in_(rec_ids))
                  .all())
        by_rec = {}
        for e in extras:
            by_rec.setdefault(e.record_id, {})[e.field_name] = e.value_text

        for r in recs:
            row = [getattr(r, t, None) for t in standard_targets]
            row += [by_rec.get(r.id, {}).get(t) for t in custom_targets]
            ws.append(row)
        last_id = recs[-1].id

    filename = f'gravitas_scrub_{job.id}_{datetime.utcnow():%Y%m%d}.xlsx'
    s3_key = storage.result_key(job.company_id, job.id, filename)
    storage.put_object(
        s3_key,
        _workbook_to_bytes(wb),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    return filename, s3_key


def generate_purchase_artifact(job):
    """Stub for purchase artifacts (purchase_jobs has no row table yet).

    Generates plausible records and uploads to Spaces, returning
    (filename, s3_key) just like the scrub variant.
    """
    headers = ['email', 'first_name', 'last_name', 'state', 'vertical', 'consent_date']
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('records')
    ws.append(headers)
    for v in (job.selected_verticals_json or []):
        n = int(v.get('records') or 0) if isinstance(v, dict) else 0
        for _ in range(n):
            ws.append([
                _rand_email(),
                random.choice(['Alex', 'Sam', 'Pat', 'Jordan', 'Casey', 'Taylor', 'Morgan']),
                random.choice(['Smith', 'Lee', 'Brown', 'Davis', 'Garcia', 'Miller', 'Wilson']),
                random.choice(['CA', 'TX', 'NY', 'FL', 'IL', 'PA', 'OH', 'GA', 'NC', 'MI']),
                v.get('vertical_name', ''),
                (datetime.utcnow() - timedelta(days=random.randint(0, 90))).isoformat(),
            ])

    filename = f'gravitas_purchase_{job.id}_{datetime.utcnow():%Y%m%d}.xlsx'
    s3_key = storage.result_key(job.company_id, job.id, filename)
    storage.put_object(
        s3_key,
        _workbook_to_bytes(wb),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    return filename, s3_key
